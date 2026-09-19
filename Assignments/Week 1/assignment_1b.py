import getpass
import json
import math
import os
import smtplib
import subprocess
import sys
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from plaid_data import get_portfolios


# ============================================================
# CONFIGURATION
# ============================================================

TARGET_PORTFOLIO = {
    "SPYM": 0.36,
    "QQQM": 0.15,
    "AVUV": 0.10,
    "VXUS": 0.10,
    "SPUU": 0.10,
    "SOXQ": 0.07,
    "GRID": 0.05,
    "MU": 0.01,
    "GOOGL": 0.01,
    "AMZN": 0.01,
    "MSFT": 0.01,
    "TSM": 0.01,
    "GEV": 0.01,
    "ORCL": 0.01,
}

TRACKING_ERROR_LIMIT = 0.005  # 0.50 percentage points
EMAIL_CONFIG_PATH = Path.home() / ".quant_python_rebalancer.json"
KEYCHAIN_SERVICE = "quant-python-rebalance"


# ============================================================
# PORTFOLIO CALCULATIONS
# ============================================================


def validate_target_weights(targets):
    total_weight = sum(targets.values())
    return round(total_weight, 4) == 1.0


def calculate_portfolio_value(portfolio):
    return sum(position["value"] for position in portfolio.values())


def calculate_current_weight(position_value, portfolio_value):
    if portfolio_value == 0:
        return 0
    return position_value / portfolio_value


def get_target_weight(ticker, targets):
    return targets.get(ticker, 0)


def calculate_weight_drift(current_weight, target_weight):
    return current_weight - target_weight


def calculate_target_value(portfolio_value, target_weight):
    return portfolio_value * target_weight


def calculate_dollar_trade(target_value, current_value):
    return target_value - current_value


def determine_action(dollar_trade, tolerance=0.005):
    if dollar_trade > tolerance:
        return "BUY"
    if dollar_trade < -tolerance:
        return "SELL"
    return "HOLD"


def calculate_share_trade(dollar_trade, price):
    if price == 0:
        return 0
    return dollar_trade / price


def build_price_lookup(portfolios):
    price_lookup = {}

    for portfolio in portfolios.values():
        for ticker, position in portfolio.items():
            if ticker != "CASH" and position.get("price", 0) > 0:
                price_lookup[ticker] = position["price"]

    return price_lookup


def build_whole_share_plan(portfolio, targets, price_lookup):
    portfolio_value = calculate_portfolio_value(portfolio)
    cash = portfolio.get("CASH", {}).get("value", 0)

    all_tickers = set(portfolio.keys()) | set(targets.keys())
    all_tickers.discard("CASH")

    plan = {}

    for ticker in all_tickers:
        if ticker in portfolio:
            current_shares = portfolio[ticker]["shares"]
            current_value = portfolio[ticker]["value"]
            price = portfolio[ticker]["price"]
        else:
            current_shares = 0
            current_value = 0
            price = price_lookup.get(ticker, 0)

        target_weight = targets.get(ticker, 0)
        target_value = portfolio_value * target_weight
        dollar_trade = target_value - current_value

        plan[ticker] = {
            "current_shares": current_shares,
            "current_value": current_value,
            "price": price,
            "target_weight": target_weight,
            "target_value": target_value,
            "dollar_trade": dollar_trade,
            "whole_share_trade": 0,
        }

    # Sell whole shares first so the proceeds are available for purchases.
    for data in plan.values():
        dollar_trade = data["dollar_trade"]
        price = data["price"]
        current_shares = data["current_shares"]

        if dollar_trade < 0 and price > 0:
            desired_sell = math.floor(abs(dollar_trade) / price)
            maximum_sell = math.floor(current_shares)
            shares_to_sell = min(desired_sell, maximum_sell)

            data["whole_share_trade"] = -shares_to_sell
            cash += shares_to_sell * price

    # Buy the largest dollar underweights first.
    buy_candidates = [
        (ticker, data)
        for ticker, data in plan.items()
        if data["dollar_trade"] > 0
    ]
    buy_candidates.sort(key=lambda item: item[1]["dollar_trade"], reverse=True)

    for _, data in buy_candidates:
        price = data["price"]
        dollar_trade = data["dollar_trade"]

        if price <= 0:
            continue

        desired_buy = math.floor(dollar_trade / price)
        affordable_buy = math.floor(cash / price)
        shares_to_buy = min(desired_buy, affordable_buy)

        data["whole_share_trade"] = shares_to_buy
        cash -= shares_to_buy * price

    post_values = {}

    for ticker, data in plan.items():
        post_shares = data["current_shares"] + data["whole_share_trade"]
        post_value = post_shares * data["price"]

        data["post_shares"] = post_shares
        data["post_value"] = post_value
        post_values[ticker] = post_value

    post_portfolio_value = sum(post_values.values()) + cash

    for data in plan.values():
        if post_portfolio_value > 0:
            post_weight = data["post_value"] / post_portfolio_value
        else:
            post_weight = 0

        weight_error = post_weight - data["target_weight"]

        data["post_weight"] = post_weight
        data["weight_error"] = weight_error

        if data["price"] <= 0 and data["target_weight"] > 0:
            data["flag"] = "NO PRICE"
        elif abs(weight_error) > TRACKING_ERROR_LIMIT:
            data["flag"] = "FLAG"
        else:
            data["flag"] = "OK"

    return plan, cash, post_portfolio_value


# ============================================================
# EMAIL CONFIGURATION / MACOS KEYCHAIN
# ============================================================


def load_email_config():
    if not EMAIL_CONFIG_PATH.exists():
        return {}

    try:
        with open(EMAIL_CONFIG_PATH, "r", encoding="utf-8") as file:
            return json.load(file)
    except (json.JSONDecodeError, OSError):
        return {}


def save_email_config(sender_email, receiver_email):
    config = {
        "sender_email": sender_email,
        "receiver_email": receiver_email,
    }

    with open(EMAIL_CONFIG_PATH, "w", encoding="utf-8") as file:
        json.dump(config, file, indent=2)

    try:
        EMAIL_CONFIG_PATH.chmod(0o600)
    except OSError:
        pass


def read_keychain_password(sender_email):
    if sys.platform != "darwin":
        return None

    result = subprocess.run(
        [
            "security",
            "find-generic-password",
            "-s",
            KEYCHAIN_SERVICE,
            "-a",
            sender_email,
            "-w",
        ],
        capture_output=True,
        text=True,
    )

    if result.returncode == 0:
        return result.stdout.strip()

    return None


def save_keychain_password(sender_email, app_password):
    if sys.platform != "darwin":
        return False

    result = subprocess.run(
        [
            "security",
            "add-generic-password",
            "-U",
            "-s",
            KEYCHAIN_SERVICE,
            "-a",
            sender_email,
            "-w",
            app_password,
        ],
        capture_output=True,
        text=True,
    )

    return result.returncode == 0


def get_email_credentials():
    config = load_email_config()

    sender_email = (
        os.environ.get("REBALANCE_EMAIL_FROM")
        or config.get("sender_email")
    )

    receiver_email = (
        os.environ.get("REBALANCE_EMAIL_TO")
        or config.get("receiver_email")
    )

    if not sender_email:
        sender_email = input("Sender Gmail address: ").strip()

    if not receiver_email:
        receiver_email = input(
            f"Recipient email [{sender_email}]: "
        ).strip() or sender_email

    # Save non-sensitive email addresses so future Terminal sessions remember them.
    save_email_config(sender_email, receiver_email)

    env_password = os.environ.get("REBALANCE_EMAIL_APP_PASSWORD")
    keychain_password = read_keychain_password(sender_email)

    if env_password:
        app_password = env_password.replace(" ", "")

        # Persist the already-authorized app password in macOS Keychain.
        if sys.platform == "darwin" and not keychain_password:
            if save_keychain_password(sender_email, app_password):
                print("Saved Gmail App Password to macOS Keychain for future runs.")

    elif keychain_password:
        app_password = keychain_password.replace(" ", "")

    else:
        print(
            "\nNo saved Gmail App Password was found. "
            "Enter it once; it will be stored in macOS Keychain."
        )
        app_password = getpass.getpass("Gmail App Password: ").replace(" ", "")

        if sys.platform == "darwin":
            if save_keychain_password(sender_email, app_password):
                print("Saved Gmail App Password to macOS Keychain.")
            else:
                print(
                    "Could not save the password to Keychain. "
                    "The current run can still continue."
                )

    return sender_email, receiver_email, app_password


# ============================================================
# EXCEL REPORT
# ============================================================


def style_header_row(sheet, row_number=1):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E1F2")

    for cell in sheet[row_number]:
        if cell.value is not None:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=thin_gray)


def auto_size_columns(sheet, minimum=10, maximum=34):
    for column_cells in sheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        sheet.column_dimensions[column_letter].width = min(
            max(max_length + 2, minimum), maximum
        )


def apply_action_formatting(sheet, action_column, start_row, end_row):
    green_fill = PatternFill("solid", fgColor="E2F0D9")
    red_fill = PatternFill("solid", fgColor="FCE4D6")
    gray_fill = PatternFill("solid", fgColor="E7E6E6")

    range_ref = f"{action_column}{start_row}:{action_column}{end_row}"

    sheet.conditional_formatting.add(
        range_ref,
        FormulaRule(
            formula=[f'${action_column}{start_row}="BUY"'],
            fill=green_fill,
        ),
    )
    sheet.conditional_formatting.add(
        range_ref,
        FormulaRule(
            formula=[f'${action_column}{start_row}="SELL"'],
            fill=red_fill,
        ),
    )
    sheet.conditional_formatting.add(
        range_ref,
        FormulaRule(
            formula=[f'${action_column}{start_row}="HOLD"'],
            fill=gray_fill,
        ),
    )


def create_excel_report(portfolios, targets, price_lookup, whole_share_results):
    workbook = Workbook()
    workbook.remove(workbook.active)

    # --------------------------
    # Dashboard
    # --------------------------
    dashboard = workbook.create_sheet("Dashboard")
    dashboard.merge_cells("A1:E1")
    dashboard["A1"] = "Portfolio Rebalance Dashboard"
    dashboard["A1"].font = Font(size=18, bold=True)
    dashboard["A1"].alignment = Alignment(horizontal="center")

    dashboard["A2"] = "Generated"
    dashboard["B2"] = datetime.now().strftime("%B %d, %Y %I:%M %p")
    dashboard["D2"] = "Target Model Valid"
    dashboard["E2"] = "YES" if validate_target_weights(targets) else "NO"

    dashboard_headers = [
        "Account",
        "Current Value",
        "Post-Trade Value",
        "Remaining Cash",
        "Flagged Positions",
    ]

    for col, header in enumerate(dashboard_headers, start=1):
        dashboard.cell(row=4, column=col, value=header)

    style_header_row(dashboard, 4)

    dashboard_start_row = 5

    for row_number, (account_name, portfolio) in enumerate(
        portfolios.items(), start=dashboard_start_row
    ):
        current_value = calculate_portfolio_value(portfolio)
        result = whole_share_results[account_name]
        flagged_positions = sum(
            1
            for data in result["plan"].values()
            if data["flag"] != "OK"
        )

        dashboard.cell(row=row_number, column=1, value=account_name)
        dashboard.cell(row=row_number, column=2, value=current_value)
        dashboard.cell(row=row_number, column=3, value=result["post_value"])
        dashboard.cell(row=row_number, column=4, value=result["cash"])
        dashboard.cell(row=row_number, column=5, value=flagged_positions)

    dashboard_end_row = dashboard_start_row + len(portfolios) - 1
    total_row = dashboard_end_row + 2

    dashboard.cell(row=total_row, column=1, value="TOTAL")
    dashboard.cell(
        row=total_row,
        column=2,
        value=f"=SUM(B{dashboard_start_row}:B{dashboard_end_row})",
    )
    dashboard.cell(
        row=total_row,
        column=3,
        value=f"=SUM(C{dashboard_start_row}:C{dashboard_end_row})",
    )
    dashboard.cell(
        row=total_row,
        column=4,
        value=f"=SUM(D{dashboard_start_row}:D{dashboard_end_row})",
    )
    dashboard.cell(
        row=total_row,
        column=5,
        value=f"=SUM(E{dashboard_start_row}:E{dashboard_end_row})",
    )

    for cell in dashboard[total_row]:
        cell.font = Font(bold=True)

    for row in dashboard.iter_rows(
        min_row=dashboard_start_row,
        max_row=total_row,
        min_col=2,
        max_col=4,
    ):
        for cell in row:
            cell.number_format = "$#,##0.00"

    chart = BarChart()
    chart.title = "Current Account Values"
    chart.y_axis.title = "Account"
    chart.x_axis.title = "Value ($)"
    chart.height = 8
    chart.width = 14

    chart_data = Reference(
        dashboard,
        min_col=2,
        min_row=4,
        max_row=dashboard_end_row,
    )
    chart_categories = Reference(
        dashboard,
        min_col=1,
        min_row=dashboard_start_row,
        max_row=dashboard_end_row,
    )

    chart.add_data(chart_data, titles_from_data=True)
    chart.set_categories(chart_categories)
    dashboard.add_chart(chart, "G4")

    dashboard.freeze_panes = "A5"
    auto_size_columns(dashboard)

    # --------------------------
    # Target Model
    # --------------------------
    target_sheet = workbook.create_sheet("Target Model")
    target_sheet.append(["Ticker", "Target Weight"])

    for ticker, weight in sorted(targets.items()):
        target_sheet.append([ticker, weight])

    target_sheet.append(["TOTAL", sum(targets.values())])
    style_header_row(target_sheet)

    for row in target_sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "0.00%"

    target_sheet.freeze_panes = "A2"
    target_sheet.auto_filter.ref = target_sheet.dimensions
    auto_size_columns(target_sheet)

    # --------------------------
    # Current Weights
    # --------------------------
    current_sheet = workbook.create_sheet("Current Weights")
    current_sheet.append(
        [
            "Account",
            "Ticker",
            "Current Value",
            "Current Weight",
            "Target Weight",
            "Drift",
        ]
    )

    for account_name, portfolio in portfolios.items():
        account_value = calculate_portfolio_value(portfolio)
        all_tickers = set(portfolio.keys()) | set(targets.keys())

        for ticker in sorted(all_tickers):
            current_value = portfolio.get(ticker, {}).get("value", 0)
            current_weight = calculate_current_weight(current_value, account_value)
            target_weight = targets.get(ticker, 0)
            drift = current_weight - target_weight

            current_sheet.append(
                [
                    account_name,
                    ticker,
                    current_value,
                    current_weight,
                    target_weight,
                    drift,
                ]
            )

    style_header_row(current_sheet)
    current_sheet.freeze_panes = "A2"
    current_sheet.auto_filter.ref = current_sheet.dimensions

    for row in current_sheet.iter_rows(min_row=2):
        row[2].number_format = "$#,##0.00"
        row[3].number_format = "0.00%"
        row[4].number_format = "0.00%"
        row[5].number_format = "+0.00%;-0.00%;0.00%"

    drift_end_row = current_sheet.max_row
    red_fill = PatternFill("solid", fgColor="FCE4D6")
    current_sheet.conditional_formatting.add(
        f"F2:F{drift_end_row}",
        CellIsRule(operator="greaterThan", formula=[str(TRACKING_ERROR_LIMIT)], fill=red_fill),
    )
    current_sheet.conditional_formatting.add(
        f"F2:F{drift_end_row}",
        CellIsRule(operator="lessThan", formula=[str(-TRACKING_ERROR_LIMIT)], fill=red_fill),
    )
    auto_size_columns(current_sheet)

    # --------------------------
    # Fractional Trades
    # --------------------------
    fractional_sheet = workbook.create_sheet("Fractional Trades")
    fractional_sheet.append(
        [
            "Account",
            "Ticker",
            "Current Value",
            "Target Value",
            "Trade Dollars",
            "Action",
            "Price",
            "Shares",
        ]
    )

    for account_name, portfolio in portfolios.items():
        account_value = calculate_portfolio_value(portfolio)
        all_tickers = set(portfolio.keys()) | set(targets.keys())
        all_tickers.discard("CASH")

        for ticker in sorted(all_tickers):
            if ticker in portfolio:
                current_value = portfolio[ticker]["value"]
                price = portfolio[ticker]["price"]
            else:
                current_value = 0
                price = price_lookup.get(ticker, 0)

            target_weight = targets.get(ticker, 0)
            target_value = account_value * target_weight
            dollar_trade = target_value - current_value
            action = determine_action(dollar_trade)
            shares = calculate_share_trade(dollar_trade, price)

            fractional_sheet.append(
                [
                    account_name,
                    ticker,
                    current_value,
                    target_value,
                    dollar_trade,
                    action,
                    price,
                    abs(shares),
                ]
            )

    style_header_row(fractional_sheet)
    fractional_sheet.freeze_panes = "A2"
    fractional_sheet.auto_filter.ref = fractional_sheet.dimensions

    for row in fractional_sheet.iter_rows(min_row=2):
        row[2].number_format = "$#,##0.00"
        row[3].number_format = "$#,##0.00"
        row[4].number_format = "$#,##0.00;[Red]-$#,##0.00"
        row[6].number_format = "$#,##0.00"
        row[7].number_format = "0.0000"

    apply_action_formatting(
        fractional_sheet,
        "F",
        2,
        fractional_sheet.max_row,
    )
    auto_size_columns(fractional_sheet)

    # --------------------------
    # Whole Share Trades
    # --------------------------
    whole_sheet = workbook.create_sheet("Whole Share Trades")
    whole_sheet.append(
        [
            "Account",
            "Ticker",
            "Current Shares",
            "Trade Shares",
            "Price",
            "Target Weight",
            "Post Shares",
            "Post Value",
            "Post Weight",
            "Weight Error",
            "Status",
        ]
    )

    for account_name, result in whole_share_results.items():
        for ticker in sorted(result["plan"]):
            data = result["plan"][ticker]
            whole_sheet.append(
                [
                    account_name,
                    ticker,
                    data["current_shares"],
                    data["whole_share_trade"],
                    data["price"],
                    data["target_weight"],
                    data["post_shares"],
                    data["post_value"],
                    data["post_weight"],
                    data["weight_error"],
                    data["flag"],
                ]
            )

    style_header_row(whole_sheet)
    whole_sheet.freeze_panes = "A2"
    whole_sheet.auto_filter.ref = whole_sheet.dimensions

    for row in whole_sheet.iter_rows(min_row=2):
        row[2].number_format = "0.0000"
        row[3].number_format = "0"
        row[4].number_format = "$#,##0.00"
        row[5].number_format = "0.00%"
        row[6].number_format = "0.0000"
        row[7].number_format = "$#,##0.00"
        row[8].number_format = "0.00%"
        row[9].number_format = "+0.00%;-0.00%;0.00%"

    green_fill = PatternFill("solid", fgColor="E2F0D9")
    red_fill = PatternFill("solid", fgColor="FCE4D6")
    yellow_fill = PatternFill("solid", fgColor="FFF2CC")

    whole_end_row = whole_sheet.max_row

    whole_sheet.conditional_formatting.add(
        f"D2:D{whole_end_row}",
        CellIsRule(operator="greaterThan", formula=["0"], fill=green_fill),
    )
    whole_sheet.conditional_formatting.add(
        f"D2:D{whole_end_row}",
        CellIsRule(operator="lessThan", formula=["0"], fill=red_fill),
    )
    whole_sheet.conditional_formatting.add(
        f"J2:J{whole_end_row}",
        CellIsRule(operator="greaterThan", formula=[str(TRACKING_ERROR_LIMIT)], fill=red_fill),
    )
    whole_sheet.conditional_formatting.add(
        f"J2:J{whole_end_row}",
        CellIsRule(operator="lessThan", formula=[str(-TRACKING_ERROR_LIMIT)], fill=red_fill),
    )
    whole_sheet.conditional_formatting.add(
        f"K2:K{whole_end_row}",
        FormulaRule(formula=['$K2="OK"'], fill=green_fill),
    )
    whole_sheet.conditional_formatting.add(
        f"K2:K{whole_end_row}",
        FormulaRule(formula=['$K2="FLAG"'], fill=red_fill),
    )
    whole_sheet.conditional_formatting.add(
        f"K2:K{whole_end_row}",
        FormulaRule(formula=['$K2="NO PRICE"'], fill=yellow_fill),
    )

    auto_size_columns(whole_sheet)

    # --------------------------
    # Workbook defaults
    # --------------------------
    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False

    reports_folder = Path(__file__).parent / "Reports"
    reports_folder.mkdir(exist_ok=True)

    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    file_path = reports_folder / f"portfolio_rebalance_{timestamp}.xlsx"
    workbook.save(file_path)

    return file_path


# ============================================================
# EMAIL DELIVERY
# ============================================================


def email_excel_report(file_path):
    sender_email, receiver_email, app_password = get_email_credentials()

    message = EmailMessage()
    message["From"] = sender_email
    message["To"] = receiver_email
    message["Subject"] = (
        "Portfolio Rebalance Report - "
        + datetime.now().strftime("%B %d, %Y")
    )

    message.set_content(
        "Attached is the latest portfolio rebalance report generated by "
        "the Quant-Python rebalancing engine.\n\n"
        "The workbook includes a dashboard, target model, current weights, "
        "fractional trades, whole-share trades, and tracking-error flags."
    )

    with open(file_path, "rb") as file:
        excel_data = file.read()

    message.add_attachment(
        excel_data,
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=file_path.name,
    )

    try:
        with smtplib.SMTP_SSL("smtp.gmail.com", 465) as smtp:
            smtp.login(sender_email, app_password)
            smtp.send_message(message)

        print(f"\nEmail sent to {receiver_email}")

    except smtplib.SMTPAuthenticationError as error:
        print(
            "\nExcel report was created, but Gmail authentication failed. "
            "Check or replace the saved App Password."
        )
        print(f"Gmail error: {error}")

    except (smtplib.SMTPException, OSError) as error:
        print("\nExcel report was created, but the email could not be sent.")
        print(f"Email error: {error}")


# ============================================================
# CONSOLE OUTPUT
# ============================================================


def print_account_summary(portfolios):
    print("\nACCOUNT SUMMARY")
    print("=" * 60)

    for account_name, portfolio in portfolios.items():
        account_value = calculate_portfolio_value(portfolio)
        print(f"{account_name:<35} ${account_value:>12,.2f}")


def print_current_weights(portfolios):
    print("\nCURRENT PORTFOLIO WEIGHTS")
    print("=" * 60)

    for account_name, portfolio in portfolios.items():
        account_value = calculate_portfolio_value(portfolio)

        print(f"\n{account_name}")
        print("-" * 60)

        for ticker, position in portfolio.items():
            current_weight = calculate_current_weight(
                position["value"],
                account_value,
            )

            print(
                f"{ticker:<8} "
                f"${position['value']:>10,.2f} "
                f"{current_weight:>9.2%}"
            )


def print_rebalance_analysis(portfolios, targets):
    print("\nALL ACCOUNT REBALANCE ANALYSIS")
    print("=" * 80)

    for account_name, portfolio in portfolios.items():
        account_value = calculate_portfolio_value(portfolio)

        print(f"\n{account_name}")
        print(f"Portfolio Value: ${account_value:,.2f}")
        print("-" * 80)
        print(
            f"{'Ticker':<8}"
            f"{'Current':>12}"
            f"{'Target':>12}"
            f"{'Drift':>12}"
        )
        print("-" * 80)

        all_tickers = set(portfolio.keys()) | set(targets.keys())

        for ticker in sorted(all_tickers):
            position_value = portfolio.get(ticker, {}).get("value", 0)
            current_weight = calculate_current_weight(position_value, account_value)
            target_weight = get_target_weight(ticker, targets)
            drift = calculate_weight_drift(current_weight, target_weight)

            print(
                f"{ticker:<8}"
                f"{current_weight:>12.2%}"
                f"{target_weight:>12.2%}"
                f"{drift:>+12.2%}"
            )


def print_fractional_trade_plan(portfolios, targets, price_lookup):
    print("\nREBALANCE TRADE PLAN - FRACTIONAL SHARES")
    print("=" * 110)

    for account_name, portfolio in portfolios.items():
        account_value = calculate_portfolio_value(portfolio)

        print(f"\n{account_name}")
        print(f"Portfolio Value: ${account_value:,.2f}")
        print("-" * 110)
        print(
            f"{'Ticker':<8}"
            f"{'Current $':>14}"
            f"{'Target $':>14}"
            f"{'Trade $':>14}"
            f"{'Action':>10}"
            f"{'Shares':>14}"
        )
        print("-" * 110)

        all_tickers = set(portfolio.keys()) | set(targets.keys())

        for ticker in sorted(all_tickers):
            if ticker == "CASH":
                continue

            if ticker in portfolio:
                current_value = portfolio[ticker]["value"]
                price = portfolio[ticker]["price"]
            else:
                current_value = 0
                price = price_lookup.get(ticker, 0)

            target_weight = get_target_weight(ticker, targets)
            target_value = calculate_target_value(account_value, target_weight)
            dollar_trade = calculate_dollar_trade(target_value, current_value)
            action = determine_action(dollar_trade)
            share_trade = calculate_share_trade(dollar_trade, price)

            print(
                f"{ticker:<8}"
                f"${current_value:>13,.2f}"
                f"${target_value:>13,.2f}"
                f"${dollar_trade:>13,.2f}"
                f"{action:>10}"
                f"{abs(share_trade):>14.4f}"
            )


def build_and_print_whole_share_results(portfolios, targets, price_lookup):
    print("\nWHOLE-SHARE REBALANCE")
    print("=" * 125)

    whole_share_results = {}

    for account_name, portfolio in portfolios.items():
        plan, remaining_cash, post_value = build_whole_share_plan(
            portfolio,
            targets,
            price_lookup,
        )

        whole_share_results[account_name] = {
            "plan": plan,
            "cash": remaining_cash,
            "post_value": post_value,
        }

        print(f"\n{account_name}")
        print(f"Post-Trade Cash: ${remaining_cash:,.2f}")
        print("-" * 125)
        print(
            f"{'Ticker':<8}"
            f"{'Trade':>10}"
            f"{'Price':>12}"
            f"{'Post Shares':>14}"
            f"{'Target':>11}"
            f"{'Post Wt':>11}"
            f"{'Error':>11}"
            f"{'Status':>10}"
        )
        print("-" * 125)

        for ticker in sorted(plan):
            data = plan[ticker]
            trade = data["whole_share_trade"]
            trade_text = f"+{trade}" if trade > 0 else str(trade)

            print(
                f"{ticker:<8}"
                f"{trade_text:>10}"
                f"${data['price']:>11,.2f}"
                f"{data['post_shares']:>14.4f}"
                f"{data['target_weight']:>11.2%}"
                f"{data['post_weight']:>11.2%}"
                f"{data['weight_error']:>+11.2%}"
                f"{data['flag']:>10}"
            )

    return whole_share_results


# ============================================================
# MAIN PROGRAM
# ============================================================

def main():
    if not validate_target_weights(TARGET_PORTFOLIO):
        raise ValueError(
            "Target portfolio weights must sum to 100%. "
            f"Current total: {sum(TARGET_PORTFOLIO.values()):.2%}"
        )

    portfolios = get_portfolios()
    price_lookup = build_price_lookup(portfolios)

    print_account_summary(portfolios)
    print_current_weights(portfolios)
    print_rebalance_analysis(portfolios, TARGET_PORTFOLIO)
    print_fractional_trade_plan(portfolios, TARGET_PORTFOLIO, price_lookup)

    whole_share_results = build_and_print_whole_share_results(
        portfolios,
        TARGET_PORTFOLIO,
        price_lookup,
    )

    report_file = create_excel_report(
        portfolios,
        TARGET_PORTFOLIO,
        price_lookup,
        whole_share_results,
    )

    print(f"\nExcel report created:\n{report_file}")
    email_excel_report(report_file)


if __name__ == "__main__":
    main()